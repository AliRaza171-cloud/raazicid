"""Long-term deprivation per district (Phase 6): where people lack schooling, safe water, toilets, solid houses,
electricity or food security over years, not just this month.

Sources publish these as files, not APIs, and often a few years old:
  PBS PSLM district survey (2019-20 is the latest district round): out-of-school children, literacy, water,
      sanitation, housing, electricity.
  PBS HIES 2024-25 (provinces only): the same subjects plus food insecurity (FIES).
  IOM DTM sheets on HDX: displaced people per district (2022 floods for Pakistan).
An admin uploads a table (CSV or Excel: a place column and a number column) and says what it measures. Displaced
people from IOM DTM also arrive by themselves through HDX HAPI (indicator "idps_present").

This is a separate layer: stored as "dep:<measure>" indicators, shown with its source and year, and NEVER mixed into
the need levels (those only use recent data). Each value gets a band from its rank among all places with that
measure: the worst quarter is "high", the next quarter "moderate", the rest "lower".
"""

import csv
import difflib
import io
import re
import uuid
from collections import defaultdict
from datetime import date
from typing import Any, Optional

from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.models import Location, NeedIndicator, Source
from app.models.enums import LocationLevel, SourceKind
from app.services.gazetteer import Place, find_places, load_places

MEASURES: dict[str, dict[str, Any]] = {
    "out_of_school": {"label": "Children out of school", "share": True, "categories": ["school_supplies", "books"],
                      "hint": "PSLM district tables or HIES 2024-25: children aged 5-16 not in school (%)"},
    "illiteracy": {"label": "People who cannot read (10+)", "share": True, "categories": ["books", "school_supplies"],
                   "hint": "PSLM table 2.14 gives the literacy rate: tick 'the table gives the opposite'"},
    "no_safe_water": {"label": "Households without improved drinking water", "share": True, "categories": ["hygiene"],
                      "hint": "PSLM / HIES water tables; if they give access to improved water, tick 'the opposite'"},
    "no_toilet": {"label": "Households without a toilet", "share": True, "categories": ["hygiene"],
                  "hint": "PSLM / HIES sanitation tables (no toilet or open defecation, %)"},
    "kutcha_housing": {"label": "Households in kutcha (mud / temporary) houses", "share": True,
                       "categories": ["blankets", "winter_support", "household", "emergency_relief"],
                       "hint": "PSLM / HIES housing tables: houses with kutcha walls or roof (%)"},
    "no_electricity": {"label": "Households without electricity", "share": True, "categories": ["household", "winter_support"],
                       "hint": "PSLM / HIES energy tables; if they give access to electricity, tick 'the opposite'"},
    "food_insecurity": {"label": "Households moderately or severely food insecure", "share": True,
                        "categories": ["food", "grocery_ration"], "hint": "HIES 2024-25 FIES tables (provinces, %)"},
    "displaced": {"label": "Displaced people", "share": False,
                  "categories": ["emergency_relief", "grocery_ration", "hygiene", "blankets", "household"],
                  "hint": "IOM DTM sheets on HDX (number of displaced people per district). Also arrives via HDX by itself"},
}
PLACE_WORDS = ("district", "admin2", "place", "area", "location", "city", "tehsil", "region", "name")
PROVINCE_WORDS = ("province", "admin1", "state")
VALUE_WORDS = ("value", "count", "number", "share", "percent", "%", "total", "population", "individuals", "rate")
SPLIT = 0.75, 0.5   # rank above the 75th percentile: high; above the median: moderate


# ---------- reading a table (pure, unit-tested) ----------

def number(text: Any, share: bool) -> Optional[float]:
    """'1,234' -> 1234; '34%', '34' or '0.34' -> 0.34 for shares. None if it is not a usable number."""
    if isinstance(text, (int, float)) and not isinstance(text, bool):
        v, pct = float(text), False
    else:
        t = str(text or "").strip().replace(",", "").replace(" ", "")
        pct = t.endswith("%")
        try:
            v = float(t.rstrip("%"))
        except ValueError:
            return None
    if share:
        if pct or v > 1:
            v = v / 100
        if not 0 <= v <= 1:
            return None
    elif v < 0:
        return None
    return v


def read_rows(raw: bytes, filename: str) -> list[list[Any]]:
    """All non-empty rows of a CSV or the first Excel sheet."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        text = raw.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = [list(r) for r in csv.reader(io.StringIO(text), dialect)]
    return [r for r in rows if any(str(c or "").strip() for c in r)]


def _hit(h: str, word: str) -> bool:
    flat = re.sub(r"[\s_\-]+", "", h)
    return h == word or word in re.split(r"[\s_\-/()]+", h) or flat.startswith(word)


def _find(head: list[str], wanted: Optional[str], words: tuple[str, ...], skip: tuple[str, ...] = ()) -> Optional[int]:
    """Column index: the one named `wanted`, else the first column matching the earliest word in `words`
    (so 'district' beats 'name'), never a column that matches a `skip` word."""
    if wanted:
        w = wanted.strip().lower()
        return next((i for i, h in enumerate(head) if h == w), next((i for i, h in enumerate(head) if w in h), None))
    for word in words:
        for i, h in enumerate(head):
            if h and _hit(h, word) and not any(_hit(h, x) for x in skip):
                return i
    return None


def parse_table(rows: list[list[Any]], share: bool, place_column: Optional[str] = None, value_column: Optional[str] = None,
                invert: bool = False) -> dict[str, Any]:
    """{"rows": [(place, province or None, value)], "problems": [...], "columns": [headers], "used": {...}}.
    The header is the first row that names a place column (or the first row). Columns are found by name, or given."""
    if not rows:
        return {"rows": [], "problems": ["The file is empty"], "columns": [], "used": {}}
    def is_header(r: list[Any]) -> bool:   # names a place column and holds no numbers (a data row has its value)
        return (_find([str(c or "").strip().lower() for c in r], place_column, PLACE_WORDS + PROVINCE_WORDS) is not None
                and all(number(c, False) is None for c in r if str(c or "").strip()))
    hi = next((i for i, r in enumerate(rows[:15]) if is_header(r)), None)
    head = [str(c or "").strip().lower() for c in rows[hi]] if hi is not None else []
    columns = [str(c or "").strip() for c in rows[hi]] if hi is not None else []
    if hi is None:   # no header: place in the first column, number in the second
        place_i, prov_i, value_i, body = 0, None, 1, rows
    else:
        body = rows[hi + 1:]
        place_i = _find(head, place_column, PLACE_WORDS, skip=PROVINCE_WORDS)
        prov_i = _find(head, None, PROVINCE_WORDS)
        if place_i is None:   # a province-only table (HIES): the province column is the place
            place_i, prov_i = prov_i, None
        if prov_i == place_i:
            prov_i = None
        value_i = _find(head, value_column, VALUE_WORDS)
        if value_i is None or value_i in (place_i, prov_i):   # the first column to the right with mostly numbers
            value_i = next((i for i in range(len(head)) if i not in (place_i, prov_i)
                            and sum(number(r[i], share) is not None for r in body if i < len(r)) >= max(1, len(body) // 2)), None)
        if value_column and _find(head, value_column, ()) is None:
            return {"rows": [], "problems": [f"No column called '{value_column}'"], "columns": columns, "used": {}}
        if place_column and _find(head, place_column, ()) is None:
            return {"rows": [], "problems": [f"No column called '{place_column}'"], "columns": columns, "used": {}}
    if value_i is None or place_i is None:
        return {"rows": [], "problems": ["Could not tell which column has the numbers: name it in 'Column with the numbers'"],
                "columns": columns, "used": {}}
    out, problems = [], []
    start = (hi + 2) if hi is not None else 1
    # PSLM tables often give Urban / Rural / Total rows per district: keep the totals only
    split_i = next((i for i in range(max((len(r) for r in body), default=0)) if i not in (place_i, value_i, prov_i)
                    and {"urban", "rural"} <= {str(r[i] or "").strip().lower() for r in body if i < len(r)}), None)
    for n, r in enumerate(body, start=start):
        if split_i is not None and str(r[split_i] if split_i < len(r) else "").strip().lower() not in ("total", "overall", "all", ""):
            continue
        if len(r) <= max(place_i, value_i):
            problems.append(f"Row {n}: missing a column")
            continue
        name = str(r[place_i] or "").strip()
        v = number(r[value_i], share)
        if not name or name.lower() in ("total", "pakistan", "overall", "all"):
            continue
        if v is None:
            problems.append(f"Row {n}: '{str(r[value_i] or '').strip()}' is not a number")
            continue
        if invert and share:
            v = round(1 - v, 4)   # e.g. literacy rate -> share who cannot read
        prov = str(r[prov_i] or "").strip() if prov_i is not None and prov_i < len(r) else None
        out.append((name, prov or None, v))
    used = {"place": columns[place_i] if hi is not None else "column 1", "value": columns[value_i] if hi is not None else "column 2",
            "province": columns[prov_i] if (hi is not None and prov_i is not None) else None}
    return {"rows": out, "problems": problems, "columns": columns, "used": used}


def clean_name(name: str) -> str:
    return re.sub(r"\b(district|division|tehsil|agency)\b", "", name, flags=re.I).strip(" ,-")


def match(name: str, places: list[Place], province: Optional[Place] = None) -> Optional[Place]:
    """Our place for a name in a table ('Kech District' -> Kech), inside the province when one is given."""
    clean = clean_name(name)
    by_id = {x.id: x for x in places}
    inside = [p for p in places if province is None or p.id == province.id or _under(p, province.id, by_id)]
    for p in inside:
        if p.name.lower() in (name.lower(), clean.lower()):
            return p
    found = find_places(clean, inside)
    if len(found) == 1:
        return found[0]
    # Spelling differences between surveys ("Qambar Shahdadkot" / "Kambar Shahdadkot"): a very close name
    close = difflib.get_close_matches(clean.lower(), [p.name.lower() for p in inside if p.level != "province"], n=1, cutoff=0.85)
    return next((p for p in inside if close and p.name.lower() == close[0]), None)


def _under(p: Place, root: uuid.UUID, by_id: dict[uuid.UUID, Place]) -> bool:
    cur, guard = p, 0
    while cur and cur.parent_id and guard < 6:
        if cur.parent_id == root:
            return True
        cur = by_id.get(cur.parent_id)
        guard += 1
    return False


def band(value: float, pool: list[float]) -> Optional[str]:
    """'high' for the worst quarter of places with this measure, 'moderate' for the next quarter, else 'lower'.
    None when fewer than 3 places have it (nothing to compare with)."""
    if len(pool) < 3:
        return None
    rank = sum(1 for x in pool if x < value) / len(pool)   # share of places doing better
    return "high" if rank >= SPLIT[0] else "moderate" if rank >= SPLIT[1] else "lower"


def display(value: float, share: bool) -> str:
    return f"{value:.0%}" if share else f"{int(round(value)):,} people"


# ---------- storing ----------

def upload_source(db: Session) -> int:
    return db.scalar(select(Source.id).where(Source.kind == SourceKind.upload))


def save_table(db: Session, measure: str, rows: list[tuple[str, Optional[str], float]], source: str, year: int,
               add_missing: bool = True) -> dict[str, Any]:
    """Replace this measure's table. Districts we do not have yet are added under their province when the table
    names the province (like the HDX districts)."""
    from app.services.sources.hdx import _district, _province
    places = load_places(db)
    sid = upload_source(db)
    db.execute(delete(NeedIndicator).where(NeedIndicator.indicator == f"dep:{measure}"))
    matched, added, unmatched, seen = 0, 0, [], set()
    for name, prov_name, v in rows:
        prov = _province(prov_name, places) if prov_name else None
        p = match(name, places, prov) or (match(name, places) if prov is None else None)
        loc_id = p.id if p else None
        if loc_id is None and prov is not None and add_missing:
            loc_id = _district(db, prov, clean_name(name), None, places).id
            added += 1
        if loc_id is None:
            unmatched.append(name)
            continue
        if loc_id in seen:
            continue   # urban / rural rows of the same district: keep the first (usually the total)
        seen.add(loc_id)
        db.add(NeedIndicator(location_id=loc_id, indicator=f"dep:{measure}", value=v, source_id=sid,
                             period_end=date(year, 12, 31), meta={"source": source[:120], "year": year, "name": name[:120]}))
        matched += 1
    db.commit()
    return {"matched": matched, "added_places": added, "unmatched": unmatched[:30], "unmatched_count": len(unmatched)}


def tables(db: Session) -> list[dict[str, Any]]:
    out = []
    for m, info in MEASURES.items():
        rows = list(db.scalars(select(NeedIndicator).where(NeedIndicator.indicator == f"dep:{m}")))
        meta = (rows[0].meta or {}) if rows else {}
        out.append({"measure": m, "label": info["label"], "hint": info["hint"], "share": info["share"], "rows": len(rows),
                    "source": meta.get("source"), "year": meta.get("year")})
    return out


# ---------- the layer for an area ----------

def load_values(db: Session) -> dict[str, dict[uuid.UUID, dict[str, Any]]]:
    """measure -> {place id: {value, source, year}}. DTM displacement from HDX fills 'displaced' where no table was
    uploaded for that place."""
    out: dict[str, dict[uuid.UUID, dict[str, Any]]] = defaultdict(dict)
    for ind in db.scalars(select(NeedIndicator).where(NeedIndicator.indicator.like("dep:%"))):
        meta = ind.meta or {}
        out[ind.indicator[4:]][ind.location_id] = {"value": ind.value, "source": meta.get("source") or "Uploaded table",
                                                   "year": meta.get("year")}
    for ind in db.scalars(select(NeedIndicator).where(NeedIndicator.indicator == "idps_present")):
        if ind.location_id not in out["displaced"]:
            ops = ", ".join((ind.meta or {}).get("operations") or [])
            out["displaced"][ind.location_id] = {"value": ind.value, "source": f"IOM DTM via HDX{' · ' + ops if ops else ''}",
                                                 "year": ind.period_end.year if ind.period_end else None}
    return {k: v for k, v in out.items() if v and k in MEASURES}


def summarise(values: dict[uuid.UUID, dict[str, Any]], inside: list[uuid.UUID], own: uuid.UUID, share: bool,
              pool: list[float]) -> Optional[dict[str, Any]]:
    """One place's value for a measure: its own figure, else the districts inside it (average of shares, sum of
    people). Pure, unit-tested."""
    if own in values:
        v = values[own]
        return {"value": v["value"], "display": display(v["value"], share), "band": band(v["value"], pool),
                "places": 1, "source": v["source"], "year": v["year"]}
    have = [values[i] for i in inside if i in values and i != own]
    if not have:
        return None
    total = sum(h["value"] for h in have)
    value = total / len(have) if share else total
    years = sorted({h["year"] for h in have if h["year"]})
    return {"value": value, "display": display(value, share) + (" (average)" if share and len(have) > 1 else ""),
            "band": band(value, pool) if share else None, "places": len(have),
            "source": have[0]["source"], "year": years[-1] if years else None}


def area_deprivation(db: Session, parent_id: Optional[uuid.UUID]) -> dict[str, Any]:
    locs = list(db.scalars(select(Location)))
    kids: dict[Optional[uuid.UUID], list[Location]] = defaultdict(list)
    for l in locs:
        kids[l.parent_id].append(l)

    def under(root: uuid.UUID) -> list[uuid.UUID]:
        out, todo = [], [root]
        while todo:
            cur = todo.pop()
            out.append(cur)
            todo += [k.id for k in kids.get(cur, [])]
        return out

    values = load_values(db)
    pools = {m: [v["value"] for v in vals.values()] for m, vals in values.items()}
    measures = [{"measure": m, "label": MEASURES[m]["label"], "share": MEASURES[m]["share"],
                 "categories": MEASURES[m]["categories"],
                 "sources": sorted({f"{v['source']} {v['year'] or ''}".strip() for v in values[m].values()})}
                for m in MEASURES if m in values]
    rows = []
    for c in sorted(kids.get(parent_id, []), key=lambda l: l.name):
        inside = under(c.id)
        vals = {m["measure"]: summarise(values[m["measure"]], inside, c.id, m["share"], pools[m["measure"]]) for m in measures}
        vals = {k: v for k, v in vals.items() if v}
        if vals:
            rows.append({"id": str(c.id), "name": c.name, "level": c.level.value, "values": vals,
                         "high": sum(1 for v in vals.values() if v["band"] == "high")})
    rows.sort(key=lambda r: (-r["high"], r["name"]))
    here = None
    if parent_id:
        inside = under(parent_id)
        here = {m["measure"]: summarise(values[m["measure"]], inside, parent_id, m["share"], pools[m["measure"]]) for m in measures}
        here = {k: v for k, v in here.items() if v}
    return {"measures": measures, "rows": rows, "here": here, "tables": tables(db)}


def facts_for(db: Session, location_id: Optional[uuid.UUID], slugs: list[str], limit: int = 6) -> list[str]:
    """Plain sentences for the need finder: the worst long-term deprivation in (or under) a place, with source and year."""
    values = load_values(db)
    if not values:
        return []
    names = {l.id: l.name for l in db.scalars(select(Location))}
    from app.services.locations import descendant_ids
    allowed = set(descendant_ids(db, location_id)) if location_id else None
    out: list[tuple[float, str]] = []
    for m, vals in values.items():
        info = MEASURES[m]
        if slugs and not set(slugs) & set(info["categories"]):
            continue
        pool = [v["value"] for v in vals.values()]
        for loc, v in vals.items():
            if allowed is not None and loc not in allowed:
                continue
            b = band(v["value"], pool)
            if b == "lower":
                continue
            rank = sum(1 for x in pool if x < v["value"]) / len(pool) if len(pool) >= 3 else 0.5
            out.append((rank, f"{names.get(loc, '?')}: {info['label'].lower()} {display(v['value'], info['share'])}"
                              f"{' (among the worst ' + str(round((1 - rank) * 100) or 1) + '% of places)' if b == 'high' else ''}"
                              f" ({v['source']}{', ' + str(v['year']) if v['year'] else ''}; long-term, not this month's need)"))
    out.sort(key=lambda x: -x[0])
    return [t for _, t in out[:limit]]
